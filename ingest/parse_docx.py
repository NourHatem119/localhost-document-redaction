"""DOCX ingest via python-docx.

DOCX is flow-based: there is no page geometry, so there are no word boxes and
no page width/height. The whole document is treated as a single logical page
with canonical text only. Embedded images are pulled from the package's
media parts and written to disk for Person B. Spatial redaction is not possible
for DOCX (no bboxes) — only the text path applies.
"""

from __future__ import annotations

from pathlib import Path
from typing import List

from docx import Document as DocxDocument

from schema import Document, ImageRef, Page

# DOCX has no real page dimensions; record 0.0 to signal "no geometry".
_NO_GEOMETRY = 0.0


def _extract_images(
    docx: "DocxDocument", doc_id: str, out_dir: Path
) -> List[ImageRef]:
    """Write embedded media parts to disk as ImageRefs (page_no=1)."""
    refs: List[ImageRef] = []
    index = 0
    for rel in docx.part.rels.values():
        if "image" not in rel.reltype:
            continue
        blob = rel.target_part.blob
        ext = rel.target_part.partname.ext.lstrip(".") or "png"
        image_id = f"{doc_id}_img{index}"
        out_path = out_dir / f"{image_id}.{ext}"
        out_path.write_bytes(blob)
        refs.append(
            ImageRef(
                image_id=image_id,
                doc_id=doc_id,
                page_no=1,
                path=str(out_path),
                bbox=None,
            )
        )
        index += 1
    return refs


def parse_docx(path: str | Path, doc_id: str, image_dir: str | Path) -> Document:
    """Parse a DOCX into a single-page Document (text only) and extract images."""
    path = Path(path)
    out_dir = Path(image_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    docx = DocxDocument(str(path))
    text = "\n".join(p.text for p in docx.paragraphs)

    # NO-GEOMETRY CONTRACT: DOCX is flow-based with no page geometry, so
    # width/height are 0.0 and word_boxes is empty. Downstream redaction must
    # treat this as a text-only source (no rectangles to draw); branch on
    # `word_boxes == []` / `bbox is None` rather than assuming page boxes exist.
    page = Page(
        page_no=1,
        width=_NO_GEOMETRY,
        height=_NO_GEOMETRY,
        text=text,
        word_boxes=[],
    )
    images = _extract_images(docx, doc_id, out_dir)

    return Document(
        doc_id=doc_id, source_path=str(path), pages=[page], images=images
    )
